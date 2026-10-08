"""Bounded file import; does not execute spreadsheets, download URLs or run notes."""
import copy
import csv
import io
import re
import zipfile
from pathlib import PurePath
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from .sequence import parse_records, SequenceError
from .store import digest, now, uid

MAX_RECORDS=1000
MAX_BASES=100000

def catalogue_rows(data,filename):
    if filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist())>30*1024*1024:
                raise SequenceError('IMPORT_LIMIT','Expanded spreadsheet exceeds 30 MB.')
        try:
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        except (KeyError,ValueError,TypeError,OSError,InvalidFileException,ParseError) as e:
            raise SequenceError('WORKBOOK_INVALID','Cannot read this workbook; export a valid values-only XLSX or CSV.',str(e)) from e
        try:
            ws=wb.active
            if ws.max_row>MAX_RECORDS+1 or ws.max_column>100:
                raise SequenceError('IMPORT_LIMIT','Catalogue limit is 1,000 rows and 100 columns.')
            values=[]
            for row in ws.iter_rows():
                if any(c.data_type=='f' for c in row):raise SequenceError('FORMULA_NOT_ALLOWED','Use a values-only catalogue; spreadsheet formulas are not evaluated.')
                values.append([str(c.value) if c.value is not None else '' for c in row])
            if not values:return []
            names=[v.strip().lower() for v in values[0]]
            if len(set(names))!=len(names):raise SequenceError('CATALOGUE_COLUMNS','Catalogue column names must be unique.')
            return [dict(zip(names,row)) for row in values[1:] if any(row)]
        finally:wb.close()
    text=data.decode('utf-8-sig')
    reader=csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:raise SequenceError('CATALOGUE_COLUMNS','Catalogue needs a header row.')
    columns=[k.strip().lower() for k in reader.fieldnames]
    if len(set(columns))!=len(columns):raise SequenceError('CATALOGUE_COLUMNS','Catalogue column names must be unique.')
    rows=[]
    for row in reader:
        if None in row:raise SequenceError('CATALOGUE_COLUMNS','A row contains more values than the header.')
        rows.append({k.strip().lower():v or '' for k,v in row.items()})
        if len(rows)>MAX_RECORDS:raise SequenceError('IMPORT_LIMIT','Catalogue limit is 1,000 rows.')
    return rows

def import_files(files,source,state,store,inspection_only=False):
    if source not in ('lab','igem','addgene','ncbi'):raise SequenceError('SOURCE_INVALID','Choose lab, iGEM, Addgene or NCBI as provenance for supplied exports.')
    parsed=[];catalogues=[];diagnostics=[];raws={}
    for filename,data in files:
        filename=PurePath(filename.replace('\\','/')).name
        if not filename or len(filename)>250:raise SequenceError('FILENAME_INVALID','Invalid input filename.')
        h=digest(data);raws[h]=data
        if filename.lower().endswith(('.csv','.xlsx')):
            if inspection_only:raise SequenceError('inspection_format','Inspection-only import accepts GenBank originals, not catalogues.')
            catalogues.append((filename,h,catalogue_rows(data,filename)));continue
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError as e:raise SequenceError('UNSUPPORTED_FORMAT','Sequence files must be UTF-8 GenBank or FASTA.') from e
        for index,r in enumerate(parse_records(text,filename,allow_parser_warnings=inspection_only)):
            if r['length']>MAX_BASES:raise SequenceError('IMPORT_LIMIT','Prototype sequence limit is 100,000 bases per record.')
            r['source'].update(kind=source,filename=filename,raw_sha256=h,imported_at=now(),access='user-supplied export; not fetched or independently authenticated')
            r['_import_key']=f'{h}:{index}:{source}'+(':inspection' if inspection_only else '');r['id']=uid('rec');parsed.append(r)
    for filename,h,rows in catalogues:
        for index,row in enumerate(rows):
            if any(str(v).lstrip().startswith(('=','+','@')) for v in row.values()):
                raise SequenceError('FORMULA_NOT_ALLOWED',f'{filename} row {index+2}: formula-like values are not accepted.')
            name=(row.get('name') or row.get('id') or '').strip()
            if not name:raise SequenceError('CATALOGUE_NAME',f'{filename} row {index+2}: add a name or id column.')
            if row.get('sequence','').strip():
                if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]{0,47}',name):raise SequenceError('CATALOGUE_NAME','Inline-sequence catalogue names must use 1–48 letters/numbers/underscore/dot/hyphen, with no spaces. Use a separate description column for display text.')
                sequence=''.join(row['sequence'].split()).upper()
                if not sequence or set(sequence)-set('ACGTRYSWKMBDHVN'):
                    raise SequenceError('CATALOGUE_SEQUENCE','The sequence cell must contain DNA IUPAC letters only, without FASTA headers or additional records.')
                parsed_row=parse_records(f'>{name}\n{sequence}\n',filename)
                if len(parsed_row)!=1:raise SequenceError('CATALOGUE_SEQUENCE','Each catalogue row must represent exactly one sequence.')
                r=parsed_row[0]
            else:
                candidates=[r for r in parsed+list(state['records'].values()) if r['name']==name or r['id']==name]
                if len(candidates)!=1:
                    raise SequenceError('CATALOGUE_MATCH',f'{filename} row {index+2}: {name} matches {len(candidates)} records; provide a unique ID or inline sequence.')
                original=candidates[0];r=copy.deepcopy(original)
                r['source']['sequence_source']=copy.deepcopy(original['source'])
                if original in parsed:parsed.remove(original)
                else:r['metadata']['parent_record_id']=original['id']
            if r['length']>MAX_BASES:raise SequenceError('IMPORT_LIMIT','Prototype sequence limit is 100,000 bases per record.')
            if row.get('topology','').strip():
                topology=row['topology'].strip().lower()
                if topology not in ('linear','circular','unknown'):raise SequenceError('TOPOLOGY_INVALID','Catalogue topology must be circular, linear or unknown.')
                if r['source'].get('format')=='genbank' and r['topology']!='unknown' and r['topology']!=topology:
                    raise SequenceError('METADATA_CONFLICT','Catalogue topology conflicts with GenBank; resolve the source files before importing.')
                r['topology']=topology
            r['metadata']['catalogue']=row
            r['source'].update(kind=source,filename=filename,raw_sha256=h,imported_at=now(),access='user-supplied catalogue')
            r['_import_key']=f'{h}:row{index}:{source}';r['id']=uid('rec');parsed.append(r)
    if len(parsed)>MAX_RECORDS:raise SequenceError('IMPORT_LIMIT','One transaction supports at most 1,000 records.')
    if not parsed:raise SequenceError('EMPTY_IMPORT','No sequence records were found.')
    accepted=[];existing=dict(state['records']);keys={r.get('_import_key'):r for r in existing.values()}
    for r in parsed:
        if r['_import_key'] in keys:
            diagnostics.append({'code':'DUPLICATE','message':f'{r["name"]}: this exact source record is already imported.','record_id':keys[r['_import_key']]['id']});continue
        conflicts=[x for x in existing.values() if x['name']==r['name']]
        if conflicts:diagnostics.append({'code':'NAME_CONFLICT','message':f'{r["name"]}: another record has this label; both versions are retained. Select exact record IDs.','record_ids':[x['id'] for x in conflicts]+[r['id']]})
        accepted.append(r);existing[r['id']]=r;keys[r['_import_key']]=r
        if r.get('metadata',{}).get('parser_warnings'):
            diagnostics.append({'code':'PARSER_WARNING','record_id':r['id'],'message':'Source retained for unchanged inspection only. Map locations are parser interpretations.',
                                'warnings':r['metadata']['parser_warnings']})
    # Parsing and conflict checks complete before any visible state mutation.
    for h,data in raws.items():store.keep_raw(data)
    state['records']=existing
    return accepted,diagnostics
